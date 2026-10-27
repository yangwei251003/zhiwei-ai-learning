from pathlib import Path
import zipfile

import pytest
from docx import Document
from pypdf import PdfWriter

from app.knowledge_base.loaders import DocumentParseError, parse_document
from app.knowledge_base import loaders


def _write_pdf(path: Path, pages: list[str]) -> None:
    """Write a small, real PDF using built-in Helvetica text streams."""
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        f"<< /Type /Pages /Kids [{' '.join(f'{3 + i * 2} 0 R' for i in range(len(pages)))}] /Count {len(pages)} >>".encode(),
    ]
    for index, text in enumerate(pages):
        page_number = 3 + index * 2
        stream_number = page_number + 1
        escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        stream = f"BT /F1 12 Tf 72 720 Td ({escaped}) Tj ET".encode()
        objects.extend(
            [
                f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 << /Type /Font /Subtype /Type1 /BaseFont /Helvetica >> >> >> /Contents {stream_number} 0 R >>".encode(),
                b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
            ]
        )

    data = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for number, obj in enumerate(objects, start=1):
        offsets.append(len(data))
        data.extend(f"{number} 0 obj\n".encode() + obj + b"\nendobj\n")
    xref = len(data)
    data.extend(f"xref\n0 {len(objects) + 1}\n".encode())
    data.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        data.extend(f"{offset:010d} 00000 n \n".encode())
    data.extend(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    path.write_bytes(data)


@pytest.mark.parametrize(("suffix", "source_type"), [(".txt", "txt"), (".MD", "md")])
def test_parse_utf8_text_and_markdown_with_bom(tmp_path: Path, suffix: str, source_type: str) -> None:
    path = tmp_path / f"notes{suffix}"
    path.write_bytes("\ufeff中文 notes".encode("utf-8"))

    segments = parse_document(path)

    assert len(segments) == 1
    assert segments[0].text == "中文 notes"
    assert segments[0].source_type == source_type
    assert (segments[0].char_start, segments[0].char_end) == (0, 8)
    assert segments[0].page is None
    assert segments[0].paragraph is None


def test_parse_docx_emits_nonempty_paragraphs_with_positions(tmp_path: Path) -> None:
    path = tmp_path / "lesson.DOCX"
    document = Document()
    document.add_paragraph("First")
    document.add_paragraph("   ")
    document.add_paragraph("第三段")
    document.save(path)

    segments = parse_document(path)

    assert [(item.text, item.paragraph) for item in segments] == [("First", 1), ("第三段", 3)]
    assert [(item.char_start, item.char_end) for item in segments] == [(0, 5), (5, 8)]
    assert all(item.source_type == "docx" and item.page is None for item in segments)


def test_parse_pdf_emits_nonempty_pages_with_positions(tmp_path: Path) -> None:
    path = tmp_path / "lesson.PDF"
    _write_pdf(path, ["Page one", "Page two"])

    segments = parse_document(path)

    assert [(item.text, item.page) for item in segments] == [("Page one", 1), ("Page two", 2)]
    assert [(item.char_start, item.char_end) for item in segments] == [(0, 8), (8, 16)]
    assert all(item.source_type == "pdf" and item.paragraph is None for item in segments)


def test_original_filename_can_supply_the_document_type(tmp_path: Path) -> None:
    path = tmp_path / "upload"
    path.write_text("content", encoding="utf-8")

    segments = parse_document(path, original_filename="NOTES.TXT", mime="text/plain")

    assert segments[0].source_type == "txt"


def test_original_filename_is_authoritative_for_document_type(tmp_path: Path) -> None:
    path = tmp_path / "temporary.txt"
    path.write_text("content", encoding="utf-8")

    with pytest.raises(DocumentParseError, match="Unsupported"):
        parse_document(path, original_filename="payload.csv")


@pytest.mark.parametrize("original_filename", ["README", ""])
def test_original_filename_without_a_supported_extension_does_not_fall_back(
    tmp_path: Path, original_filename: str
) -> None:
    path = tmp_path / "temporary.txt"
    path.write_text("content", encoding="utf-8")

    with pytest.raises(DocumentParseError, match="Unsupported"):
        parse_document(path, original_filename=original_filename, mime="text/plain")


@pytest.mark.parametrize(
    ("filename", "contents", "message"),
    [
        ("empty.txt", b"", "empty"),
        ("spaces.md", b" \n\t", "empty"),
        ("bad.txt", b"\xff\xfe", "UTF-8"),
        ("data.csv", b"a,b", "Unsupported"),
        ("bad.docx", b"not a zip", "DOCX"),
        ("bad.pdf", b"not a pdf", "PDF"),
    ],
)
def test_invalid_or_unsupported_documents_raise_clear_errors(
    tmp_path: Path, filename: str, contents: bytes, message: str
) -> None:
    path = tmp_path / filename
    path.write_bytes(contents)

    with pytest.raises(DocumentParseError, match=message):
        parse_document(path)


def test_pdf_without_extractable_text_is_rejected(tmp_path: Path) -> None:
    path = tmp_path / "blank.pdf"
    _write_pdf(path, [""])

    with pytest.raises(DocumentParseError, match="extractable text"):
        parse_document(path)


def test_encrypted_pdf_is_rejected(tmp_path: Path) -> None:
    path = tmp_path / "protected.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    writer.encrypt("secret")
    with path.open("wb") as output:
        writer.write(output)

    with pytest.raises(DocumentParseError, match="Encrypted PDF"):
        parse_document(path)


@pytest.mark.parametrize(
    ("filename", "contents"),
    [
        ("renamed.pdf", b"plain text"),
        ("renamed.docx", b"%PDF-1.4\n"),
        ("renamed.txt", b"%PDF-1.4\n"),
        ("renamed.md", b"PK\x03\x04binary"),
        ("bom-renamed.txt", b"\xef\xbb\xbf%PDF-1.4\n"),
        ("nul.txt", b"hello\x00world"),
    ],
)
def test_content_signatures_reject_renamed_or_binary_files(
    tmp_path: Path, filename: str, contents: bytes
) -> None:
    path = tmp_path / filename
    path.write_bytes(contents)

    with pytest.raises(DocumentParseError, match="content does not match"):
        parse_document(path)


def test_mime_must_match_declared_extension(tmp_path: Path) -> None:
    path = tmp_path / "notes.txt"
    path.write_text("content", encoding="utf-8")

    with pytest.raises(DocumentParseError, match="MIME type does not match"):
        parse_document(path, mime="application/pdf")


@pytest.mark.parametrize("mime", ["application/octet-stream", "application/x-unknown"])
def test_extensionless_path_with_unknown_mime_is_unsupported(
    tmp_path: Path, mime: str
) -> None:
    path = tmp_path / "upload"
    _write_pdf(path, ["valid PDF text"])

    with pytest.raises(DocumentParseError, match="Unsupported document type"):
        parse_document(path, mime=mime)


def test_file_size_budget_is_checked_before_reading(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    path = tmp_path / "large.txt"
    path.write_text("12345", encoding="utf-8")
    monkeypatch.setattr(loaders, "MAX_FILE_BYTES", 4)

    with pytest.raises(DocumentParseError, match="size limit"):
        parse_document(path)


def test_pdf_page_and_extracted_text_budgets(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    path = tmp_path / "pages.pdf"
    _write_pdf(path, ["1234", "5678"])
    monkeypatch.setattr(loaders, "MAX_PDF_PAGES", 1)

    with pytest.raises(DocumentParseError, match="page limit"):
        parse_document(path)

    monkeypatch.setattr(loaders, "MAX_PDF_PAGES", 2)
    monkeypatch.setattr(loaders, "MAX_PDF_EXTRACTED_CHARS", 7)
    with pytest.raises(DocumentParseError, match="text limit"):
        parse_document(path)


def _write_docx_like_zip(path: Path, members: dict[str, bytes]) -> None:
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for name, contents in members.items():
            archive.writestr(name, contents)


def test_docx_requires_core_package_members(tmp_path: Path) -> None:
    path = tmp_path / "incomplete.docx"
    _write_docx_like_zip(path, {"other.xml": b"<root/>"})

    with pytest.raises(DocumentParseError, match="required package members"):
        parse_document(path)


def test_docx_rejects_a_valid_zip_with_a_binary_prefix(tmp_path: Path) -> None:
    path = tmp_path / "prefixed.docx"
    document = Document()
    document.add_paragraph("text")
    document.save(path)
    path.write_bytes(b"MZ\x00\x00" + path.read_bytes())

    with pytest.raises(DocumentParseError, match="content does not match"):
        parse_document(path)


def test_docx_zip_bomb_budgets(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    path = tmp_path / "budget.docx"
    members = {
        "[Content_Types].xml": b"<Types/>",
        "word/document.xml": b"x" * 200,
        "extra.xml": b"x",
    }
    _write_docx_like_zip(path, members)

    monkeypatch.setattr(loaders, "MAX_DOCX_MEMBERS", 2)
    with pytest.raises(DocumentParseError, match="member limit"):
        parse_document(path)

    monkeypatch.setattr(loaders, "MAX_DOCX_MEMBERS", 10)
    monkeypatch.setattr(loaders, "MAX_DOCX_MEMBER_BYTES", 100)
    with pytest.raises(DocumentParseError, match="member size limit"):
        parse_document(path)

    monkeypatch.setattr(loaders, "MAX_DOCX_MEMBER_BYTES", 1_000)
    monkeypatch.setattr(loaders, "MAX_DOCX_UNCOMPRESSED_BYTES", 100)
    with pytest.raises(DocumentParseError, match="uncompressed size limit"):
        parse_document(path)

    monkeypatch.setattr(loaders, "MAX_DOCX_UNCOMPRESSED_BYTES", 10_000)
    monkeypatch.setattr(loaders, "MAX_DOCX_COMPRESSION_RATIO", 2)
    with pytest.raises(DocumentParseError, match="compression ratio limit"):
        parse_document(path)


def test_errors_do_not_disclose_paths_and_files_are_deletable(tmp_path: Path) -> None:
    missing = tmp_path / "private-folder" / "secret.txt"
    with pytest.raises(DocumentParseError) as missing_error:
        parse_document(missing)
    assert str(missing) not in str(missing_error.value)

    invalid = tmp_path / "bad.pdf"
    invalid.write_bytes(b"not-pdf")
    with pytest.raises(DocumentParseError) as invalid_error:
        parse_document(invalid)
    assert str(invalid) not in str(invalid_error.value)
    invalid.unlink()
    assert not invalid.exists()

    valid = tmp_path / "valid.pdf"
    _write_pdf(valid, ["text"])
    parse_document(valid)
    valid.unlink()
    assert not valid.exists()

    valid_docx = tmp_path / "valid.docx"
    document = Document()
    document.add_paragraph("text")
    document.save(valid_docx)
    parse_document(valid_docx)
    valid_docx.unlink()
    assert not valid_docx.exists()

    invalid_docx = tmp_path / "invalid.docx"
    invalid_docx.write_bytes(b"not-docx")
    with pytest.raises(DocumentParseError):
        parse_document(invalid_docx)
    invalid_docx.unlink()
    assert not invalid_docx.exists()


@pytest.mark.parametrize("source_type", ["pdf", "docx"])
def test_parser_receives_a_file_object_that_is_closed_after_success(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, source_type: str
) -> None:
    path = tmp_path / f"valid.{source_type}"
    streams = []
    if source_type == "pdf":
        _write_pdf(path, ["text"])
        original_parser = loaders.PdfReader

        def tracked_parser(stream):
            streams.append(stream)
            return original_parser(stream)

        monkeypatch.setattr(loaders, "PdfReader", tracked_parser)
    else:
        document = Document()
        document.add_paragraph("text")
        document.save(path)
        original_parser = loaders.Document

        def tracked_parser(stream):
            streams.append(stream)
            return original_parser(stream)

        monkeypatch.setattr(loaders, "Document", tracked_parser)

    parse_document(path)

    assert streams and streams[0].closed


@pytest.mark.parametrize("source_type", ["pdf", "docx"])
def test_parser_file_object_is_closed_after_parser_exception(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, source_type: str
) -> None:
    path = tmp_path / f"valid.{source_type}"
    if source_type == "pdf":
        _write_pdf(path, ["text"])
    else:
        document = Document()
        document.add_paragraph("text")
        document.save(path)
    streams = []

    def failing_parser(stream):
        streams.append(stream)
        raise RuntimeError("parser detail")

    monkeypatch.setattr(loaders, "PdfReader" if source_type == "pdf" else "Document", failing_parser)

    with pytest.raises(DocumentParseError) as error:
        parse_document(path)

    assert "parser detail" not in str(error.value)
    assert streams and streams[0].closed
