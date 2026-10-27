from __future__ import annotations

from io import BytesIO
from zipfile import ZipFile

import pytest
from docx import Document

from app.services.documents import (
    DocumentParseError,
    ParsedSection,
    chunk_sections,
    parse_document,
    validate_document_upload,
)


def test_parse_text_and_markdown_documents() -> None:
    text = parse_document("notes.txt", "第一段\n\n第二段".encode())
    markdown = parse_document("lesson.md", "# 标题\n\n正文".encode())

    assert [section.text for section in text] == ["第一段\n\n第二段"]
    assert text[0].metadata == {"source": "notes.txt"}
    assert markdown[0].text == "# 标题\n\n正文"


def test_parse_docx_keeps_non_empty_paragraphs() -> None:
    document = Document()
    document.add_paragraph("第一段")
    document.add_paragraph("")
    document.add_paragraph("第二段")
    payload = BytesIO()
    document.save(payload)

    sections = parse_document("lesson.docx", payload.getvalue())

    assert [section.text for section in sections] == ["第一段", "第二段"]
    assert [section.metadata["paragraph"] for section in sections] == [1, 3]


def test_parse_pdf_keeps_page_numbers(monkeypatch) -> None:
    class Page:
        def __init__(self, text: str | None) -> None:
            self.text = text

        def extract_text(self) -> str | None:
            return self.text

    class Reader:
        def __init__(self, _: BytesIO) -> None:
            self.pages = [Page("第一页"), Page(None), Page("第三页")]

    monkeypatch.setattr("app.services.documents.PdfReader", Reader)

    sections = parse_document("guide.pdf", b"pdf payload")

    assert [section.text for section in sections] == ["第一页", "第三页"]
    assert [section.metadata["page"] for section in sections] == [1, 3]


def test_parse_document_rejects_unsupported_or_broken_content() -> None:
    with pytest.raises(DocumentParseError, match="不支持"):
        parse_document("archive.zip", b"content")
    with pytest.raises(DocumentParseError, match="UTF-8"):
        parse_document("broken.txt", b"\xff")


def test_chunk_sections_enforces_size_overlap_and_metadata() -> None:
    sections = [ParsedSection("甲" * 12, {"source": "lesson.txt", "page": 2})]

    chunks = chunk_sections(sections, chunk_size=8, overlap=3)

    assert [chunk.text for chunk in chunks] == ["甲" * 8, "甲" * 7]
    assert chunks[0].text[-3:] == chunks[1].text[:3]
    assert [chunk.metadata["chunk_index"] for chunk in chunks] == [0, 1]
    assert all(chunk.metadata["page"] == 2 for chunk in chunks)


def test_chunk_sections_rejects_invalid_limits_and_skips_blank_sections() -> None:
    with pytest.raises(ValueError, match="overlap"):
        chunk_sections([ParsedSection("内容", {})], chunk_size=5, overlap=5)

    assert chunk_sections([ParsedSection("   ", {"source": "blank.txt"})]) == []


def test_upload_validation_rejects_empty_mismatched_and_forged_files() -> None:
    with pytest.raises(DocumentParseError, match="不能为空"):
        validate_document_upload("empty.txt", "text/plain", b"")
    with pytest.raises(DocumentParseError, match="类型"):
        validate_document_upload("lesson.txt", "application/pdf", b"plain text")
    with pytest.raises(DocumentParseError, match="PDF.*签名"):
        validate_document_upload("lesson.pdf", "application/pdf", b"not a pdf")
    with pytest.raises(DocumentParseError, match="DOCX.*签名"):
        validate_document_upload(
            "lesson.docx",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            b"not a zip",
        )


def test_upload_validation_rejects_docx_with_excessive_zip_entries() -> None:
    payload = BytesIO()
    with ZipFile(payload, "w") as archive:
        archive.writestr("[Content_Types].xml", "types")
        archive.writestr("word/document.xml", "document")
        for index in range(6):
            archive.writestr(f"word/item-{index}.xml", "x")

    with pytest.raises(DocumentParseError, match="条目"):
        validate_document_upload(
            "lesson.docx",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            payload.getvalue(),
            max_docx_entries=5,
        )
